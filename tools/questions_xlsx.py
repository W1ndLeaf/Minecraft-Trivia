"""questions.xlsx <-> question dicts.

Sheet "Questions", one row per question:
  ID | Pool (Normal/Hard) | Topic | Question | A | B | C | D | Correct (A-D) | Explanation | Source | Version
Rows with an empty Question are ignored, so you can leave gaps. C and D may be empty (2- or 3-option questions).
"""
import sys

try:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:  # pragma: no cover
    sys.exit("openpyxl is missing:  pip install openpyxl")

HEADERS = ["ID", "Pool", "Topic", "Question", "A", "B", "C", "D", "Correct", "Explanation", "Source", "Version"]
WIDTHS = [9, 9, 12, 60, 28, 28, 28, 28, 9, 60, 40, 12]
POOL_NAMES = {"Normal": "N", "Hard": "H"}


def _s(v):
    return "" if v is None else str(v).strip()


def read_questions(path):
    """Returns a list of dicts {id, pool: 'N'|'H', topic, q, options, correct, explain, source, version, row}.
    Raises SystemExit with every problem listed (row numbers refer to the sheet)."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if "Questions" not in wb.sheetnames:
        sys.exit(f"{path}: no sheet named 'Questions'")
    ws = wb["Questions"]
    rows = ws.iter_rows(values_only=True)
    header = [_s(c) for c in next(rows)]
    col = {}
    for name in HEADERS:
        if name not in header:
            sys.exit(f"{path}: header row must contain {HEADERS}, missing {name!r}")
        col[name] = header.index(name)
    out, problems, seen_ids = [], [], set()
    for rno, row in enumerate(rows, 2):
        row = list(row) + [None] * len(HEADERS)
        get = lambda name: _s(row[col[name]])  # noqa: E731
        if not get("Question"):
            continue
        pool = get("Pool").capitalize()
        if pool not in POOL_NAMES:
            problems.append(f"row {rno}: Pool must be Normal or Hard (got {get('Pool')!r})")
            continue
        options = [get(k) for k in ("A", "B", "C", "D")]
        while options and not options[-1]:
            options.pop()
        if any(not o for o in options) or not 2 <= len(options) <= 4:
            problems.append(f"row {rno}: fill A and B (C, D optional, no gaps)")
            continue
        letter = get("Correct").upper()[:1]
        if letter not in "ABCD"[:len(options)] or not letter:
            problems.append(f"row {rno}: Correct must be one of {'/'.join('ABCD'[:len(options)])}")
            continue
        qid = get("ID") or f"row{rno}"
        if qid in seen_ids:
            problems.append(f"row {rno}: duplicate ID {qid}")
            continue
        seen_ids.add(qid)
        out.append({"id": qid, "pool": POOL_NAMES[pool], "topic": get("Topic"), "q": get("Question"),
                    "options": options, "correct": "ABCD".index(letter), "explain": get("Explanation"),
                    "source": get("Source"), "version": get("Version"), "row": rno})
    if problems:
        sys.exit(f"{path}: fix these rows first:\n  " + "\n  ".join(problems))
    if not any(q["pool"] == "N" for q in out) or not any(q["pool"] == "H" for q in out):
        sys.exit(f"{path}: need at least one Normal and one Hard question")
    return out


def write_questions(path, questions):
    """Creates a fresh workbook with the questions, dropdowns and a help sheet."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Questions"
    ws.append(HEADERS)
    for c in range(1, len(HEADERS) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2F5597")
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(c)].width = WIDTHS[c - 1]
    for q in questions:
        opts = list(q["options"]) + [""] * (4 - len(q["options"]))
        ws.append([q["id"], "Normal" if q["pool"] == "N" else "Hard", q.get("topic", ""), q["q"],
                   opts[0], opts[1], opts[2], opts[3], "ABCD"[q["correct"]], q.get("explain", ""),
                   q.get("source", ""), q.get("version", "")])
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{max(len(questions) + 1, 2)}"
    dv_pool = DataValidation(type="list", formula1='"Normal,Hard"', allow_blank=True, showErrorMessage=True,
                             errorTitle="Pool", error="Normal or Hard")
    dv_ans = DataValidation(type="list", formula1='"A,B,C,D"', allow_blank=True, showErrorMessage=True,
                            errorTitle="Correct", error="A, B, C or D")
    ws.add_data_validation(dv_pool)
    ws.add_data_validation(dv_ans)
    dv_pool.add("B2:B2000")
    dv_ans.add("I2:I2000")

    h = wb.create_sheet("How to edit")
    h.column_dimensions["A"].width = 110
    for line in [
        "HOW TO EDIT THE QUESTIONS",
        "",
        "- One row = one question. Add rows at the bottom, delete rows you do not want. Empty rows are ignored.",
        "- Pool: Normal or Hard (dropdown). Normal = the easier pool, Hard = deep-cut knowledge.",
        "- Question: the text shown in game (keep it to about 2 lines; plain letters, no emoji).",
        "- A, B, C, D: the answers. C and D may be empty for 2- or 3-answer questions (no gaps: fill A, B, then C, then D).",
        "- Correct: the letter of the right answer (dropdown A-D).",
        "- ID: any short unique text (N-077, H-080 ...). Topic / Explanation / Source / Version are notes for you, not shown in game.",
        "",
        "AFTER EDITING",
        "  1. Save this file (keep the name questions.xlsx, keep the sheet name Questions).",
        "  2. Run:  python tools/build.py      (needs Python 3 and:  pip install openpyxl)",
        "  3. The rebuilt packs are in dist/. The build stops and tells you the row number if something is wrong.",
        "",
        "One special question is worth keeping: 'Who made this modpack?' with B = W1ndLeaf.",
    ]:
        h.append([line])
    h["A1"].font = Font(bold=True, size=13)
    h["A10"].font = Font(bold=True)
    wb.save(path)


if __name__ == "__main__":  # quick check:  python questions_xlsx.py ../questions.xlsx
    qs = read_questions(sys.argv[1])
    print(len(qs), "questions,", sum(q["pool"] == "N" for q in qs), "Normal,", sum(q["pool"] == "H" for q in qs), "Hard")
